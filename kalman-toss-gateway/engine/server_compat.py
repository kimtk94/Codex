from __future__ import annotations

import errno
import os
import shutil
import subprocess
import zipfile
from pathlib import Path
from typing import Any

_ORIGINAL_COPY2 = shutil.copy2
_COPY2_PATCHED = False
_CRYPTO_BRIDGE_PATCHED = False


def _data_root() -> Path:
    return Path(os.environ.get("KALMAN_DATA_ROOT", "/opt/kalman/data")).expanduser().resolve()


def _under_data_root(path: str | os.PathLike[str]) -> bool:
    try:
        Path(path).expanduser().resolve().relative_to(_data_root())
        return True
    except (OSError, ValueError):
        return False


def _rclone_safe_copy2(
    src: str | os.PathLike[str],
    dst: str | os.PathLike[str],
    *,
    follow_symlinks: bool = True,
):
    try:
        return _ORIGINAL_COPY2(src, dst, follow_symlinks=follow_symlinks)
    except OSError as exc:
        # rclone/FUSE can expose normal file bytes while listxattr(2) returns
        # EIO. shutil.copy2() copies bytes first, then copystat() probes xattrs.
        # For the configured data root only, keep the bytes and skip metadata.
        if exc.errno != errno.EIO or not (
            _under_data_root(src) or _under_data_root(dst)
        ):
            raise
        return shutil.copyfile(src, dst, follow_symlinks=follow_symlinks)


def _install_rclone_copy_compat() -> None:
    global _COPY2_PATCHED
    if _COPY2_PATCHED:
        return
    shutil.copy2 = _rclone_safe_copy2
    _COPY2_PATCHED = True


def _is_real_xlsx(path: Path) -> bool:
    try:
        return path.is_file() and path.stat().st_size > 0 and zipfile.is_zipfile(path)
    except OSError:
        return False


def _mounted_crypto_sheet() -> Path | None:
    explicit = os.environ.get("KALMAN_CRYPTO_SHEET_XLSX", "").strip()
    if explicit:
        p = Path(explicit).expanduser()
        if p.exists():
            return p

    root = _data_root()
    matches = sorted(
        root.glob("Kalman Upbit KRW History*.xlsx"),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    return matches[0] if matches else None


def _export_crypto_sheet(source_hint: Path) -> Path:
    remote = os.environ.get("KALMAN_RCLONE_REMOTE", "gdrive:").strip() or "gdrive:"
    if not remote.endswith(":"):
        remote += ":"

    config = Path(
        os.environ.get("KALMAN_RCLONE_CONFIG", "/etc/rclone/rclone.conf")
    ).expanduser()
    cache = Path(
        os.environ.get("KALMAN_CRYPTO_SHEET_CACHE", "/tmp/kalman_upbit_crypto.xlsx")
    ).expanduser()
    cache.parent.mkdir(parents=True, exist_ok=True)

    tmp = cache.with_name(f".{cache.name}.tmp-{os.getpid()}")
    tmp.unlink(missing_ok=True)
    remote_source = f"{remote}{source_hint.name}"

    cmd = [
        "rclone",
        "copyto",
        remote_source,
        str(tmp),
        f"--config={config}",
    ]
    try:
        subprocess.run(
            cmd,
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            timeout=120,
        )
        if not _is_real_xlsx(tmp):
            raise RuntimeError(
                f"rclone export did not produce a valid XLSX: {remote_source}"
            )
        os.replace(tmp, cache)
    finally:
        tmp.unlink(missing_ok=True)

    return cache


def _resolve_crypto_xlsx() -> Path | None:
    source_hint = _mounted_crypto_sheet()
    if source_hint is None:
        return None

    # A Google native Sheet can appear through rclone mount as a zero-byte
    # placeholder with an .xlsx suffix. Use it directly only if it is a real
    # XLSX ZIP; otherwise request a direct rclone export to a local cache.
    if _is_real_xlsx(source_hint):
        return source_hint
    return _export_crypto_sheet(source_hint)


class _LocalWorksheet:
    def __init__(self, workbook_path: Path, title: str):
        self.workbook_path = workbook_path
        self.title = title

    @staticmethod
    def _normalize_rows(rows: list[tuple[Any, ...]]) -> list[list[Any]]:
        normalized = [
            ["" if value is None else value for value in row]
            for row in rows
        ]
        # gspread omits trailing completely-empty rows and trailing empty cells.
        while normalized and not any(value != "" for value in normalized[-1]):
            normalized.pop()
        for row in normalized:
            while row and row[-1] == "":
                row.pop()
        return normalized

    def get(self, range_name: str) -> list[list[Any]]:
        from openpyxl import load_workbook
        from openpyxl.utils.cell import range_boundaries

        min_col, min_row, max_col, max_row = range_boundaries(range_name)
        wb = load_workbook(self.workbook_path, read_only=True, data_only=True)
        try:
            if self.title not in wb.sheetnames:
                raise KeyError(
                    f"Worksheet {self.title!r} not found in {self.workbook_path}"
                )
            ws = wb[self.title]
            rows = list(
                ws.iter_rows(
                    min_row=min_row,
                    max_row=max_row,
                    min_col=min_col,
                    max_col=max_col,
                    values_only=True,
                )
            )
            return self._normalize_rows(rows)
        finally:
            wb.close()

    def get_all_values(self) -> list[list[Any]]:
        from openpyxl import load_workbook

        wb = load_workbook(self.workbook_path, read_only=True, data_only=True)
        try:
            if self.title not in wb.sheetnames:
                raise KeyError(
                    f"Worksheet {self.title!r} not found in {self.workbook_path}"
                )
            ws = wb[self.title]
            rows = list(
                ws.iter_rows(
                    min_row=1,
                    max_row=ws.max_row,
                    min_col=1,
                    max_col=ws.max_column,
                    values_only=True,
                )
            )
            return self._normalize_rows(rows)
        finally:
            wb.close()


class _LocalSpreadsheet:
    def __init__(self, workbook_path: Path):
        self.workbook_path = workbook_path

    def worksheet(self, title: str) -> _LocalWorksheet:
        return _LocalWorksheet(self.workbook_path, title)


class _LocalGspreadClient:
    def __init__(self, workbook_path: Path):
        self.workbook_path = workbook_path

    def open_by_key(self, _spreadsheet_id: str) -> _LocalSpreadsheet:
        return _LocalSpreadsheet(self.workbook_path)


class _LocalCredentials:
    """Sentinel credentials used only by the local-XLSX gspread bridge."""


def _install_crypto_xlsx_bridge() -> None:
    global _CRYPTO_BRIDGE_PATCHED
    if _CRYPTO_BRIDGE_PATCHED:
        return

    run_mode = os.environ.get("RUN_MODE", "").upper()
    if run_mode not in {"CRYPTO", "CRYPTO_GLOBAL", "FULL"}:
        return

    adc = os.environ.get("GOOGLE_APPLICATION_CREDENTIALS", "").strip()
    if adc and Path(adc).expanduser().is_file():
        return

    workbook_path = _resolve_crypto_xlsx()
    if workbook_path is None:
        return

    import google.auth
    import gspread

    def _local_default(*_args, **_kwargs):
        return _LocalCredentials(), None

    def _local_authorize(_credentials, *_args, **_kwargs):
        return _LocalGspreadClient(workbook_path)

    google.auth.default = _local_default
    gspread.authorize = _local_authorize
    _CRYPTO_BRIDGE_PATCHED = True
    print(f"[server] CRYPTO Google Sheet bridge: exported XLSX {workbook_path}")


def install_server_runtime_compat() -> None:
    _install_rclone_copy_compat()
    _install_crypto_xlsx_bridge()
