from __future__ import annotations

import errno
from pathlib import Path

import pytest

from engine import server_compat


def test_rclone_safe_copy2_falls_back_on_eio_under_data_root(tmp_path, monkeypatch):
    root = tmp_path / "drive"
    root.mkdir()
    src = root / "source.txt"
    dst = root / "dest.txt"
    src.write_text("kalman", encoding="utf-8")
    monkeypatch.setenv("KALMAN_DATA_ROOT", str(root))

    def fail_copy2(*_args, **_kwargs):
        raise OSError(errno.EIO, "Input/output error")

    monkeypatch.setattr(server_compat, "_ORIGINAL_COPY2", fail_copy2)
    result = server_compat._rclone_safe_copy2(src, dst)

    assert Path(result) == dst
    assert dst.read_text(encoding="utf-8") == "kalman"


def test_rclone_safe_copy2_does_not_hide_unrelated_eio(tmp_path, monkeypatch):
    root = tmp_path / "drive"
    root.mkdir()
    src = tmp_path / "outside.txt"
    dst = tmp_path / "outside-copy.txt"
    src.write_text("kalman", encoding="utf-8")
    monkeypatch.setenv("KALMAN_DATA_ROOT", str(root))

    def fail_copy2(*_args, **_kwargs):
        raise OSError(errno.EIO, "Input/output error")

    monkeypatch.setattr(server_compat, "_ORIGINAL_COPY2", fail_copy2)

    with pytest.raises(OSError) as exc:
        server_compat._rclone_safe_copy2(src, dst)

    assert exc.value.errno == errno.EIO


def _write_test_workbook(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Overview"
    ws["A1"] = "Kalman"
    ws["B8"] = "kalman-model-web-v8-hourly-entry-filter"
    ws["A12"] = "Market"
    ws["B12"] = "Last Close"
    ws["A13"] = "KRW-BTC"
    ws["B13"] = 109_867_000
    wb.save(path)
    wb.close()


def test_local_crypto_xlsx_reads_gspread_style_range(tmp_path):
    path = tmp_path / "Kalman Upbit KRW History 2026-08-22.xlsx"
    _write_test_workbook(path)

    values = server_compat._LocalWorksheet(path, "Overview").get("A1:B13")

    assert values[0][0] == "Kalman"
    assert values[7][1] == "kalman-model-web-v8-hourly-entry-filter"
    assert values[12] == ["KRW-BTC", 109_867_000]


def test_resolve_crypto_xlsx_exports_zero_byte_rclone_placeholder(
    tmp_path, monkeypatch
):
    root = tmp_path / "drive"
    root.mkdir()
    placeholder = root / "Kalman Upbit KRW History 2026-08-22.xlsx"
    placeholder.touch()
    cache = tmp_path / "cache" / "crypto.xlsx"

    monkeypatch.setenv("KALMAN_DATA_ROOT", str(root))
    monkeypatch.setenv("KALMAN_CRYPTO_SHEET_CACHE", str(cache))
    monkeypatch.setenv("KALMAN_RCLONE_REMOTE", "gdrive:")
    monkeypatch.setenv("KALMAN_RCLONE_CONFIG", "/etc/rclone/rclone.conf")
    monkeypatch.delenv("KALMAN_CRYPTO_SHEET_XLSX", raising=False)

    def fake_run(cmd, **kwargs):
        assert cmd[0:3] == [
            "rclone",
            "copyto",
            "gdrive:Kalman Upbit KRW History 2026-08-22.xlsx",
        ]
        exported = Path(cmd[3])
        exported.parent.mkdir(parents=True, exist_ok=True)
        _write_test_workbook(exported)

        class Result:
            returncode = 0

        return Result()

    monkeypatch.setattr(server_compat.subprocess, "run", fake_run)

    resolved = server_compat._resolve_crypto_xlsx()

    assert resolved == cache
    assert server_compat._is_real_xlsx(cache)
    assert (
        server_compat._LocalWorksheet(cache, "Overview").get("A13:B13")[0]
        == ["KRW-BTC", 109_867_000]
    )


def test_local_crypto_xlsx_get_all_values_trims_empty_tail(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "Kalman Upbit KRW History.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sources_Audit"
    ws["A1"] = "source"
    ws["B1"] = "status"
    ws["A2"] = "Upbit"
    ws["B2"] = "OK"
    wb.save(path)
    wb.close()

    values = server_compat._LocalWorksheet(path, "Sources_Audit").get_all_values()

    assert values == [["source", "status"], ["Upbit", "OK"]]


def test_kr_source_lineage_overlay_freezes_existing_live_raw(tmp_path, monkeypatch):
    import pandas as pd

    root = tmp_path / "drive"
    finance = root / "Finance_KR"
    cache = finance / "marcap_cache"
    cache.mkdir(parents=True)

    live = pd.DataFrame(
        {
            "Date": pd.to_datetime(["2026-09-16", "2026-09-16"]),
            "Code": ["000001", "000002"],
            "Open": [100.0, 200.0],
            "High": [110.0, 210.0],
            "Low": [90.0, 190.0],
            "Close": [101.0, 202.0],
            "Volume": [1000.0, 2000.0],
            "Amount": [101000.0, 404000.0],
            "Marcap": [1_000_000.0, 2_000_000.0],
            "Stocks": [10_000.0, 20_000.0],
        }
    )
    live.to_parquet(finance / "kr_hgb_live_features_v0_4_latest.parquet", index=False)

    source = live.copy()
    source.loc[source["Code"].eq("000001"), "Close"] = 103.0
    source.loc[:, "Volume"] = [9999.0, 8888.0]
    source.loc[:, "Amount"] = [999999.0, 888888.0]
    source = pd.concat(
        [
            source,
            pd.DataFrame(
                {
                    "Date": [pd.Timestamp("2026-09-17")],
                    "Code": ["000001"],
                    "Open": [102.0],
                    "High": [112.0],
                    "Low": [92.0],
                    "Close": [104.0],
                    "Volume": [1200.0],
                    "Amount": [124800.0],
                    "Marcap": [1_010_000.0],
                    "Stocks": [10_000.0],
                }
            ),
        ],
        ignore_index=True,
    )

    monkeypatch.setenv("KALMAN_DATA_ROOT", str(root))
    monkeypatch.setenv("RUN_MODE", "KR_GLOBAL")

    import pandas as real_pd
    original = real_pd.read_parquet

    def wrapper(path, *args, **kwargs):
        return original(path, *args, **kwargs)

    wrapper.__kalman_original__ = original
    monkeypatch.setattr(real_pd, "read_parquet", wrapper)

    out = server_compat._overlay_frozen_kr_live_raw(source)

    frozen = out.loc[
        out["Date"].eq(pd.Timestamp("2026-09-16"))
        & out["Code"].eq("000001")
    ].iloc[0]
    future = out.loc[
        out["Date"].eq(pd.Timestamp("2026-09-17"))
        & out["Code"].eq("000001")
    ].iloc[0]

    assert frozen["Close"] == 101.0
    assert frozen["Volume"] == 1000.0
    assert frozen["Amount"] == 101000.0
    assert future["Close"] == 104.0
