#!/usr/bin/env python3
"""Read-only Salesforce subscription mail ingestion to isolated Google Sheets staging.

Does not modify existing imported reports, decisions, financial facts, or Salesforce.
Each source is promoted to its own staging tab only after validation.
"""
import csv
import datetime as dt
import email
import hashlib
import imaplib
import io
import json
import os
import re
import sys
import time
import urllib.parse
from email import policy
from email.utils import parsedate_to_datetime
from pathlib import Path

from google.auth.transport.requests import AuthorizedSession
from google.oauth2 import service_account
from openpyxl import load_workbook

from quotation_email_ingest import TableParser

SHEET = os.getenv("SFDC_UNIFIED_SHEET_ID", "1Fo0Tc6PQ3Ek0-33lmegScM_z8ST-ShUYQ0K3kD4PqnI")
KST = dt.timezone(dt.timedelta(hours=9))
SCHEMAS = {
    "Leads": {"required": {"SFDC Lead ID", "Lead ID"}, "min_rows": 50, "key": "SFDC Lead ID", "tab": "SFDC Email Leads Raw"},
    "Opportunities": {"required": {"Opportunity ID", "Stage"}, "min_rows": 100, "key": "Opportunity ID", "tab": "SFDC Email Opportunities Raw"},
    "Quotation": {"required": {"SFDC Quotation ID", "Quotation Product ID", "Opportunity ID"}, "min_rows": 100, "key": "SFDC Quotation ID", "tab": "SFDC Email Quotation Raw"},
}
AUDIT_TAB = "SFDC Email Ingest Audit"
AUDIT_HEADER = ["Entity", "Source Received KST", "Ingest Checked KST", "Rows", "Unique IDs", "Duplicate IDs", "Missing IDs", "Content SHA256", "State", "Source Kind"]

def normalize(name):
    return re.sub(r"\s+", " ", str(name or "").replace("\ufeff", "").strip()).casefold()

def matrix(blob, filename):
    if len(blob) > 12_000_000:
        raise ValueError("ATTACHMENT_TOO_LARGE")
    if filename.lower().endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
        try:
            values = [["" if x is None else str(x).strip() for x in row] for row in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        probe = blob[:1024].lower()
        if b"charset=euc-kr" in probe or b"ks_c_5601" in probe:
            content = blob.decode("euc-kr")
        elif b"charset=cp949" in probe:
            content = blob.decode("cp949")
        else:
            try:
                content = blob.decode("utf-8-sig")
            except UnicodeDecodeError:
                content = blob.decode("cp949")
        if "<table" in content[:20000].lower() or "<tr" in content[:20000].lower():
            parser = TableParser()
            parser.feed(content)
            values = parser.rows
        else:
            values = list(csv.reader(io.StringIO(content)))
    values = [[str(v or "").strip() for v in row] for row in values if any(str(v or "").strip() for v in row)]
    if len(values) < 2:
        raise ValueError("EMPTY_REPORT")
    # Salesforce HTML reports can contain a title/banner before the real table header.
    for pos, row in enumerate(values[:12]):
        names = {normalize(v) for v in row}
        if any({normalize(x) for x in conf["required"]}.issubset(names) for conf in SCHEMAS.values()):
            columns = [str(v).strip() for v in row]
            if len(columns) != len({normalize(v) for v in columns}) or any(not c for c in columns):
                raise ValueError("DUPLICATE_OR_EMPTY_HEADER")
            rows = [(r + [""] * len(columns))[:len(columns)] for r in values[pos + 1:]]
            return columns, rows
    raise ValueError("UNRECOGNIZED_REPORT_SCHEMA")

def classify(columns):
    found = {normalize(c) for c in columns}
    hits = [kind for kind, cfg in SCHEMAS.items() if {normalize(x) for x in cfg["required"]}.issubset(found)]
    if len(hits) != 1:
        raise ValueError("AMBIGUOUS_REPORT_SCHEMA")
    return hits[0]

def validate(kind, columns, rows):
    cfg = SCHEMAS[kind]
    if len(rows) < cfg["min_rows"]:
        raise ValueError("TOO_FEW_ROWS")
    key_index = next(i for i, h in enumerate(columns) if normalize(h) == normalize(cfg["key"]))
    keys = [r[key_index] for r in rows]
    missing = sum(not value for value in keys)
    if missing:
        raise ValueError("MISSING_REQUIRED_IDS")
    unique = len(set(keys))
    duplicates = len(keys) - unique
    if kind != "Quotation" and duplicates:
        raise ValueError("DUPLICATE_PRIMARY_IDS")
    return {"rows": len(rows), "unique": unique, "duplicates": duplicates, "missing": missing}

def digest(columns, rows):
    payload = json.dumps([columns, rows], ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()

def request(session, method, suffix, **kwargs):
    root = "https://sheets.googleapis.com/v4/spreadsheets/" + SHEET
    for retry in range(4):
        result = session.request(method, root + suffix, timeout=60, **kwargs)
        if result.status_code in (429, 500, 502, 503, 504) and retry < 3:
            time.sleep(2 ** retry)
            continue
        if not result.ok:
            raise RuntimeError("SHEETS_HTTP_" + str(result.status_code))
        return result.json() if result.content else {}

def a1(tab, range_):
    return urllib.parse.quote("'" + tab.replace("'", "''") + "'!" + range_, safe="")

def column_name(index):
    value = ""
    while index:
        index, rem = divmod(index - 1, 26)
        value = chr(65 + rem) + value
    return value

def write_verified(session, selected, dry_run=False):
    meta = request(session, "GET", "?fields=sheets.properties(title,sheetId,gridProperties)")
    existing = {s["properties"]["title"]: s["properties"] for s in meta.get("sheets", [])}
    changes = []
    for kind, item in selected.items():
        tab = SCHEMAS[kind]["tab"]
        rows = item["rows"]
        cols = len(item["headers"])
        count = len(rows) + 1
        if tab not in existing:
            changes.append({"addSheet": {"properties": {"title": tab, "gridProperties": {"rowCount": max(100, count), "columnCount": max(26, cols)}}}})
        else:
            props = existing[tab]
            grid = props.get("gridProperties", {})
            if grid.get("rowCount", 0) < count or grid.get("columnCount", 0) < cols:
                changes.append({"updateSheetProperties": {"properties": {"sheetId": props["sheetId"], "gridProperties": {"rowCount": max(grid.get("rowCount", 0), count), "columnCount": max(grid.get("columnCount", 0), cols)}}, "fields": "gridProperties.rowCount,gridProperties.columnCount"}})
    if AUDIT_TAB not in existing:
        changes.append({"addSheet": {"properties": {"title": AUDIT_TAB, "gridProperties": {"rowCount": 100, "columnCount": 12}}}})
    if dry_run:
        return
    if changes:
        request(session, "POST", ":batchUpdate", json={"requests": changes})
    # Only staging tabs may be replaced. No existing business sheet is touched.
    for kind, item in selected.items():
        tab = SCHEMAS[kind]["tab"]
        old_rows = existing.get(tab, {}).get("gridProperties", {}).get("rowCount", 0)
        old_cols = existing.get(tab, {}).get("gridProperties", {}).get("columnCount", 0)
        # One atomic values write including a blank tail; no pre-clear window.
        end_rows = max(old_rows, len(item["rows"]) + 1)
        end_cols = max(old_cols, len(item["headers"]))
        replacement = [item["headers"]] + item["rows"]
        replacement = [row + [""] * (end_cols - len(row)) for row in replacement]
        replacement.extend([[""] * end_cols for _ in range(end_rows - len(replacement))])
        request(session, "POST", "/values:batchUpdate", json={"valueInputOption": "RAW", "data": [{"range": "'" + tab + "'!A1:" + column_name(end_cols) + str(end_rows), "values": replacement}]})
    checked = dt.datetime.now(KST).strftime("%Y-%m-%d %H:%M:%S KST")
    audit = [AUDIT_HEADER]
    for kind in SCHEMAS:
        item = selected.get(kind)
        if item is None:
            audit.append([kind, "", checked, 0, 0, 0, 0, "", "NO_VALID_REPORT_PRESERVED", "QQ IMAP"])
        else:
            stats = item["stats"]
            audit.append([kind, item["stamp"], checked, stats["rows"], stats["unique"], stats["duplicates"], stats["missing"], item["sha"], "STAGED_NOT_PROMOTED", "QQ IMAP"])
    request(session, "PUT", "/values/" + a1(AUDIT_TAB, "A1") + "?valueInputOption=RAW", json={"values": audit})

def collect(im, days=4):
    since = (dt.datetime.now(dt.timezone.utc) - dt.timedelta(days=days)).strftime("%d-%b-%Y")
    ok, result = im.uid("SEARCH", None, "SINCE", since)
    if ok != "OK":
        raise RuntimeError("IMAP_SEARCH_FAILED")
    ok2, targeted = im.uid("SEARCH", None, "SINCE", since, "SUBJECT", '"Report results"')
    uids = set((result[0] or b"").split()[-750:])
    if ok2 == "OK":
        uids.update((targeted[0] or b"").split())
    selected = {}
    for uid in sorted(uids, key=lambda b: int(b)):
        ok, records = im.uid("FETCH", uid, "(BODY.PEEK[])")
        if ok != "OK":
            continue
        raw = next((x[1] for x in records if isinstance(x, tuple)), None)
        if not raw:
            continue
        msg = email.message_from_bytes(raw, policy=policy.default)
        subject = str(msg.get("subject", "")).lower()
        sender = str(msg.get("from", "")).lower()
        if not (re.search(r"@(?:[a-z0-9-]+\.)*salesforce\.com\b", sender) or "salesforce" in subject or "report results" in subject):
            continue
        try:
            when = parsedate_to_datetime(str(msg["date"]))
            if when is None or when.tzinfo is None:
                continue
            when = when.astimezone(dt.timezone.utc)
        except (ValueError, TypeError, OverflowError):
            continue
        for part in msg.walk():
            filename = str(part.get_filename() or "")
            if not filename.lower().endswith((".xls", ".xlsx", ".csv")):
                continue
            payload = part.get_payload(decode=True)
            if not payload or len(payload) > 12_000_000:
                continue
            try:
                columns, rows = matrix(payload, filename)
                kind = classify(columns)
                stats = validate(kind, columns, rows)
            except (UnicodeError, ValueError, TypeError, IndexError):
                continue
            old = selected.get(kind)
            if old is not None and when <= old["when"]:
                continue
            selected[kind] = {"when": when, "stamp": when.astimezone(KST).strftime("%Y-%m-%d %H:%M:%S KST"), "headers": columns, "rows": rows, "stats": stats, "sha": digest(columns, rows)}
    return selected

def main():
    dry = os.getenv("SFDC_REPORT_DRY_RUN", "true").lower() != "false"
    with imaplib.IMAP4_SSL("imap.exmail.qq.com", 993, timeout=45) as im:
        im.login(os.environ["CSS_QQ_MAIL_ADDRESS"], os.environ["CSS_QQ_MAIL_AUTH_CODE"])
        im.select("INBOX", readonly=True)
        selected = collect(im)
    # Print counts and checksums; never log customer contents.
    print("SFDC_3REPORT_AUDIT=" + json.dumps({k: {"received": v["stamp"], **v["stats"], "sha256": v["sha"]} for k, v in selected.items()}, sort_keys=True))
    if dry:
        print("SFDC_3REPORT_DRY_RUN=true")
        return
    if not selected:
        print("SFDC_3REPORT_NO_VALID_ATTACHMENT; previous staging preserved")
        return
    cred = service_account.Credentials.from_service_account_info(json.loads(os.environ["GCP_SERVICE_ACCOUNT_JSON"]), scopes=["https://www.googleapis.com/auth/spreadsheets"])
    write_verified(AuthorizedSession(cred), selected)
    print("SFDC_3REPORT_STAGED=" + ",".join(sorted(selected)))

if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print("SFDC_3REPORT_FAILED=" + type(exc).__name__, file=sys.stderr)
        raise
