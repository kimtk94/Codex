#!/usr/bin/env python3
"""P1 isolated validation: n8n webhook contract and MarkItDown DOCX conversion."""
import argparse
import json
import tempfile
import threading
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path
from urllib.request import Request, urlopen


def ci_payload(run_id, conclusion, url):
    if conclusion not in {"success", "failure", "cancelled"}:
        raise ValueError("unsupported conclusion")
    return {
        "event": "kalman.ci.completed",
        "source": "github_actions",
        "run_id": str(run_id),
        "conclusion": conclusion,
        "run_url": url,
        "idempotency_key": f"github_actions:{run_id}:{conclusion}",
    }


def post_webhook(url, payload):
    req = Request(url, data=json.dumps(payload).encode("utf-8"),
                  headers={"Content-Type": "application/json"}, method="POST")
    with urlopen(req, timeout=4) as response:
        if not 200 <= response.status < 300:
            raise RuntimeError(f"unexpected status {response.status}")
        return response.status


def n8n_mock_smoke():
    received = []

    class Handler(BaseHTTPRequestHandler):
        def do_POST(self):
            received.append(json.loads(self.rfile.read(int(self.headers["Content-Length"]))))
            self.send_response(200)
            self.end_headers()
            self.wfile.write(b"ok")

        def log_message(self, *_args):
            pass

    server = HTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        for status in ("success", "failure", "cancelled"):
            payload = ci_payload("123", status, "https://github.com/kimtk94/Codex/actions/runs/123")
            assert post_webhook(f"http://127.0.0.1:{server.server_port}/webhook/kalman-ci", payload) == 200
        assert len(received) == 3
        assert {x["conclusion"] for x in received} == {"success", "failure", "cancelled"}
        assert len({x["idempotency_key"] for x in received}) == 3
        return {"status": "PASS", "scope": "local_mock_only", "events": len(received)}
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def markitdown_smoke():
    from docx import Document
    from markitdown import MarkItDown

    with tempfile.TemporaryDirectory() as d:
        path = Path(d) / "kalman_research.docx"
        doc = Document()
        doc.add_heading("Kalman Ischemic Stroke Research", 1)
        doc.add_paragraph("허혈성 뇌졸중 다중오믹스 연구")
        doc.add_heading("Methods", 2)
        doc.add_paragraph("MR and colocalization validation.")
        doc.save(path)
        converted = MarkItDown().convert(str(path)).text_content
        for phrase in ("Kalman Ischemic Stroke Research", "허혈성 뇌졸중", "MR and colocalization"):
            if phrase not in converted:
                raise AssertionError(f"Conversion lost expected content: {phrase}")
        return {"status": "PASS", "format": "docx", "unicode": "PASS", "characters": len(converted)}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--test", choices=["n8n", "markitdown", "all"], default="all")
    parser.add_argument("--output")
    args = parser.parse_args()
    result = {}
    if args.test in ("n8n", "all"):
        result["n8n"] = n8n_mock_smoke()
    if args.test in ("markitdown", "all"):
        result["markitdown"] = markitdown_smoke()
    evidence = json.dumps(result, ensure_ascii=False, indent=2)
    print(evidence)
    if args.output:
        Path(args.output).write_text(evidence + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
