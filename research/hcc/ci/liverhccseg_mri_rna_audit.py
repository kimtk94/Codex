#!/usr/bin/env python3
"""LiverHccSeg v1.1 MRI ↔ TCGA-LIHC tumor RNA metadata feasibility.

Downloads ONLY two small public Zenodo XLSX metadata tables (not 2.6GB images).
The LiverHccSeg cohort is MRI, *not* CT. A workbook case barcode is not proof
that the specific tumor-mask NIfTI exists or aligns with an MR sequence.
"""
import argparse
import csv
import hashlib
import json
import re
import time
import urllib.request
from collections import Counter, defaultdict
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

from gdc_tumor_lineage_audit import classify_file
from tcga_rna_overlap_audit import fetch_gdc_files, normalize_id

RECORD = "8179129"
BASE = f"https://zenodo.org/records/{RECORD}/files/"
FILES = {
    "LiverHccSeg_MetaData.xlsx": "37806f09955aa198ab8e50e0e2929da7",
    "segm_metrics.xlsx": "e3bf1618f0c0fc65b2c9d18a90a36717",
}
BARCODE = re.compile(r"TCGA[-_][A-Z0-9]{2}[-_][A-Z0-9]{4}", re.I)
COLUMNS = [
    "patient_id", "mri_metadata_workbooks",
    "mri_metadata_cell_mentions",
    "tumor_roi_nii_status", "hcc_tumor_annotation_metrics_present", "mri_sequence_phase_status",
    "primary_tumor_star_counts_files", "gdc_biospecimen_status",
    "mri_rna_case_barcode_overlap", "verified_mri_tumor_roi_rna_pair",
]


def fetch_workbook(name, expected_md5):
    import openpyxl
    url = BASE + name + "?download=1"
    last = None
    for attempt in range(3):
        try:
            request = urllib.request.Request(
                url, headers={"User-Agent": "HCC-Public-MRI-RNA-Feasibility/1.0"})
            with urllib.request.urlopen(request, timeout=60) as response:
                data = response.read(5 * 1024 * 1024 + 1)
            if len(data) > 5 * 1024 * 1024:
                raise ValueError("Unexpected metadata XLSX size above 5 MiB")
            digest = hashlib.md5(data).hexdigest()
            if expected_md5 and digest.lower() != expected_md5.lower():
                raise ValueError(f"Zenodo checksum mismatch for {name}: {digest}")
            return openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True), digest, len(data)
        except Exception as exc:
            last = exc
            if attempt < 2:
                time.sleep(2 ** attempt)
    raise RuntimeError(f"Cannot obtain verified metadata file {name}: {last}")


def scan_book(book, filename):
    """Scan explicit TCGA barcodes; do not derive patient ID from row number."""
    patients = defaultdict(int)
    tabs = []
    tumor_metric_ids = set()
    for ws in book.worksheets:
        rows = ws.iter_rows(values_only=True)
        first = next(rows, tuple())
        tabs.append({"sheet": ws.title, "rows": ws.max_row, "cols": ws.max_column,
                     "headers": [str(x)[:100] if x is not None else "" for x in first[:30]]})
        dice_column = next((i for i, h in enumerate(first) if str(h).lower().strip() == "dice"), None)
        for row in rows:
            seen_on_row = set()
            for cell in row:
                if cell is None:
                    continue
                for match in BARCODE.findall(str(cell)):
                    patient = normalize_id(match.replace("_", "-"))
                    if patient:
                        seen_on_row.add(patient)
            for patient in seen_on_row:
                patients[patient] += 1
            if filename == "segm_metrics.xlsx" and ws.title.strip().lower() == "hcc":
                dice = row[dice_column] if dice_column is not None and dice_column < len(row) else None
                if isinstance(dice, (int, float)) and 0 <= dice <= 1:
                    tumor_metric_ids.update(seen_on_row)
    return patients, tabs, tumor_metric_ids


def RNA_index(files):
    """Only file-linked primary solid tumor sample code 01, not case-level status."""
    by_patient = defaultdict(set)
    for item in files:
        if item.get("data_type") != "Gene Expression Quantification":
            continue
        if (item.get("analysis") or {}).get("workflow_type") != "STAR - Counts":
            continue
        fid = item.get("file_id") or item.get("id")
        if not fid:
            continue
        for case in item.get("cases") or []:
            pid = normalize_id(case.get("submitter_id"))
            if pid and classify_file(item, pid)[1] == "PRIMARY_TUMOR_FILE_BARCODE_LINKED":
                by_patient[pid].add(fid)
    return by_patient


def combine(books, rna_index, tumor_metric_ids=None):
    tumor_metric_ids = set(tumor_metric_ids or [])
    mentions = defaultdict(dict)
    for filename, counts in books.items():
        for patient, n in counts.items():
            mentions[patient][filename] = n
    rows = []
    for patient, origins in sorted(mentions.items()):
        tumor_files = len(rna_index.get(patient, set()))
        rows.append({
            "patient_id": patient,
            "mri_metadata_workbooks": "|".join(sorted(origins)),
            "mri_metadata_cell_mentions": sum(origins.values()),
            "tumor_roi_nii_status": "NOT_VERIFIED",
            "hcc_tumor_annotation_metrics_present": "YES" if patient in tumor_metric_ids else "NO",
            "mri_sequence_phase_status": "NOT_VERIFIED",
            "primary_tumor_star_counts_files": tumor_files,
            "gdc_biospecimen_status": ("FILE_LINKED_PRIMARY_TUMOR" if tumor_files else
                                      "NO_FILE_LINKED_PRIMARY_TUMOR"),
            "mri_rna_case_barcode_overlap": "YES" if tumor_files else "NO",
            "verified_mri_tumor_roi_rna_pair": "NOT_ASSESSED",
        })
    return rows


def offline_test():
    examples = ["TCGA-DD-A115-01A-11R", "MR study TCGA-G3-A25S"]
    assert normalize_id(BARCODE.findall(examples[0])[0]) == "TCGA-DD-A115"
    assert normalize_id(BARCODE.findall(examples[1])[0]) == "TCGA-G3-A25S"
    rows = combine({"metadata.xlsx": {"TCGA-DD-A115": 3,
                                     "TCGA-G3-A25S": 2},
                    "metric.xlsx": {"TCGA-DD-A115": 1}},
                   {"TCGA-DD-A115": {"f1"}})
    assert len(rows) == 2
    assert sum(x["mri_rna_case_barcode_overlap"] == "YES" for x in rows) == 1
    assert all(x["verified_mri_tumor_roi_rna_pair"] == "NOT_ASSESSED" for x in rows)
    print("HCC_LIVERHCCSEG_MRI_RNA_OFFLINE_TEST_PASS")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", default="artifacts/hcc_tcia")
    parser.add_argument("--offline-test", action="store_true")
    args = parser.parse_args()
    if args.offline_test:
        offline_test()
        return

    out = Path(args.output)
    out.mkdir(exist_ok=True, parents=True)
    inventory, books, hcc_metric_ids = {}, {}, set()
    try:
        for filename, checksum in FILES.items():
            wb, digest, nbytes = fetch_workbook(filename, checksum)
            counts, sheets, tumor_metric_ids = scan_book(wb, filename)
            if filename == "segm_metrics.xlsx":
                hcc_metric_ids = tumor_metric_ids
            wb.close()
            inventory[filename] = {
                "md5_verified": digest, "bytes": nbytes,
                "tcga_patient_ids_detected": len(counts), "sheets": sheets,
            }
            books[filename] = counts
        all_ids = set().union(*(set(x) for x in books.values()))
        if not all_ids:
            raise ValueError("No TCGA barcodes found in either workbook; inspect workbook format")
        rna = RNA_index(fetch_gdc_files())
        patients = combine(books, rna, hcc_metric_ids)
        overlap = sum(x["mri_rna_case_barcode_overlap"] == "YES" for x in patients)
        with (out / "tcga_liverhccseg_mri_rna_candidates.csv").open(
                "w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=COLUMNS)
            writer.writeheader()
            writer.writerows(patients)
        summary = {
            "status": "PASS_METADATA_MRI_RNA_SCREEN_ONLY",
            "generated_utc": datetime.now(timezone.utc).isoformat(),
            "zenodo_record": RECORD,
            "metadata_workbook_tcga_patients": len(all_ids),
            "mri_metadata_rna_primary_tumor_file_linked_patients": overlap,
            "unique_patients_with_hcc_segmentation_metrics": len(hcc_metric_ids),
            "hcc_segmentation_metrics_and_rna_linked_patients": sum(
                p["hcc_tumor_annotation_metrics_present"] == "YES"
                and p["mri_rna_case_barcode_overlap"] == "YES" for p in patients),
            "mri_tumor_mask_nii_verified_patients": "NOT_ASSESSED",
            "qualified_mri_tumor_rna_pairs": "NOT_ASSESSED",
            "workbooks": inventory,
            "notes": [
                "The publication describes 14 HCC tumor-segmented MRI patients. Workbook metadata IDs are NOT verified tumor masks.",
                "MRI and CT are distinct modality cohorts; no cross-modality substitution.",
                "Raw MRI NIfTI archive not downloaded; confirm actual tumor NIfTI and phase correspondence before extraction.",
                "Even a tumor MRI mask and RNA file on the same TCGA case require temporal/clinical eligibility checks.",
            ],
        }
    except Exception as exc:
        summary = {"status": "MRI_METADATA_AUDIT_FAILED",
                   "error_class": type(exc).__name__, "error": str(exc)[:500],
                   "workbooks": inventory,
                   "qualified_mri_tumor_rna_pairs": "NOT_ASSESSED"}
    (out / "tcga_liverhccseg_mri_rna_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False))
    if summary["status"] == "MRI_METADATA_AUDIT_FAILED":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
